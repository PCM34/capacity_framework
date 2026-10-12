"""Regression test for the colored Excel export.

Writes against the same frozen fixture workbook as test_capacity_model.py,
then reads the .xlsx back and checks every cell's fill against the
[0.85, 1) -> yellow, [1, inf) -> red thresholds -- on both the Workcenter
Utilization sheet and the Workcenter Hours Load sheet, since both must be
colored identically, driven by utilization.
"""
from pathlib import Path

import openpyxl
import pytest

from hub.tools.capacity_model import run_capacity_model
from hub.tools.capacity_model_export import RED_FILL, YELLOW_FILL, write_capacity_model_workbook
from hub.tools.excel_io import load_workbook

FIXTURES_DIR = Path(__file__).resolve().parent / 'fixtures'
WORKBOOK = FIXTURES_DIR / 'example_workbook.xlsx'

NO_FILL_RGB = '00000000'


@pytest.fixture(scope='module')
def result() -> dict:
    sheets = load_workbook(WORKBOOK)
    return run_capacity_model(sheets)


def test_cell_colors_match_utilization_thresholds(result: dict, tmp_path) -> None:
    out_path = tmp_path / 'export.xlsx'
    write_capacity_model_workbook(result, out_path)

    wb = openpyxl.load_workbook(out_path)
    util_ws = wb['Workcenter Utilization']
    hours_ws = wb['Workcenter Hours Load']

    util_pivot = result['util_pivot']
    n_rows, n_cols = util_pivot.shape
    assert n_rows > 0 and n_cols > 0, 'fixture should have at least one workcenter/month'

    checked_yellow = checked_red = checked_none = 0
    for r in range(n_rows):
        for c in range(n_cols):
            value = util_pivot.iat[r, c]
            excel_row, excel_col = r + 2, c + 2
            util_rgb = util_ws.cell(row=excel_row, column=excel_col).fill.start_color.rgb
            hours_rgb = hours_ws.cell(row=excel_row, column=excel_col).fill.start_color.rgb

            if value >= 1:
                assert util_rgb == RED_FILL.start_color.rgb
                checked_red += 1
            elif value >= 0.85:
                assert util_rgb == YELLOW_FILL.start_color.rgb
                checked_yellow += 1
            else:
                assert util_rgb == NO_FILL_RGB
                checked_none += 1

            # Hours sheet must match the utilization sheet cell-for-cell.
            assert hours_rgb == util_rgb

    # Sanity check the fixture actually exercises all three cases, so this test
    # can't silently pass by never hitting the yellow/red branches.
    assert checked_yellow > 0
    assert checked_red > 0
    assert checked_none > 0
