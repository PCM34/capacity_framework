"""Writing capacity model results to a formatted Excel workbook.

Pure I/O/formatting logic, no UI imports.
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import PatternFill

# Excel's own standard "Highlight Cell Rules" colors (Yellow Fill / Light Red Fill),
# chosen so text stays readable rather than a pure #FFFF00/#FF0000 block.
YELLOW_FILL = PatternFill(start_color='FFFFEB9C', end_color='FFFFEB9C', fill_type='solid')
RED_FILL = PatternFill(start_color='FFFFC7CE', end_color='FFFFC7CE', fill_type='solid')

UTIL_WARNING_THRESHOLD = 0.85  # [0.85, 1) -> yellow
UTIL_OVER_THRESHOLD = 1.0      # [1, inf)  -> red


def _fill_for_utilization(value: float) -> PatternFill | None:
    if value >= UTIL_OVER_THRESHOLD:
        return RED_FILL
    if value >= UTIL_WARNING_THRESHOLD:
        return YELLOW_FILL
    return None


def write_capacity_model_workbook(result: dict, path: str | Path) -> None:
    """Write Raw Capacity Output, Workcenter Utilization, and Workcenter Hours Load to
    one .xlsx, one sheet each.

    Utilization cells in [0.85, 1) are highlighted yellow and [1, inf) red. The Hours
    table gets the identical highlighting at the same (Workcenter, Month) cell -- the
    color is always driven by the utilization value, so both views flag the same
    overloaded periods consistently.
    """
    raw_output = result['raw_output']
    util_pivot = result['util_pivot']
    hours_pivot = result['hours_pivot']

    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        raw_output.to_excel(writer, sheet_name='Raw Capacity Output', index=False)
        util_pivot.reset_index().to_excel(writer, sheet_name='Workcenter Utilization', index=False)
        hours_pivot.reset_index().to_excel(writer, sheet_name='Workcenter Hours Load', index=False)

        util_sheet = writer.sheets['Workcenter Utilization']
        hours_sheet = writer.sheets['Workcenter Hours Load']

        n_rows, n_cols = util_pivot.shape
        for r in range(n_rows):
            for c in range(n_cols):
                fill = _fill_for_utilization(util_pivot.iat[r, c])
                if fill is None:
                    continue
                excel_row = r + 2  # +1 for the header row, +1 for openpyxl's 1-indexing
                excel_col = c + 2  # +1 for the Workcenter index column, +1 for 1-indexing
                util_sheet.cell(row=excel_row, column=excel_col).fill = fill
                hours_sheet.cell(row=excel_row, column=excel_col).fill = fill
